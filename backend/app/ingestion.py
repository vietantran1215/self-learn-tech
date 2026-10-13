import base64
from langsmith import traceable
from pathlib import Path
from sqlalchemy import delete
from .db_models import DocumentRecord, ChunkRecord
import pymupdf as fitz
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation
from .config import settings
from .models import llm, embed, content_text
from .store import db, event, uid

async def describe_image(raw: bytes, mime: str):
    if settings.demo_mode:
        raise ValueError("Image understanding requires live mode and a vision-capable model")
    result = await llm().ainvoke([{"role":"user","content":[
        {"type":"text","text":"Transcribe this image and describe diagrams, charts, labels and relationships precisely. Do not invent unreadable content."},
        {"type":"image_url","image_url":{"url":f"data:{mime};base64,{base64.b64encode(raw).decode()}"}}
    ]}])
    return content_text(result)

async def extract(path: Path):
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return [(1,path.read_text(encoding="utf-8"),"text")]
    if suffix in {".png", ".jpg", ".jpeg"}:
        return [(1, await describe_image(path.read_bytes(),"image/png" if suffix==".png" else "image/jpeg"),"image")]
    if suffix == ".pptx":
        deck = Presentation(path)
        pages=[]
        for i,slide in enumerate(deck.slides,1):
            pages.append((i,"\n".join(s.text for s in slide.shapes if s.has_text_frame),"text"))
            for shape in slide.shapes:
                if shape.shape_type == 13 and not settings.demo_mode:
                    pages.append((i,await describe_image(shape.image.blob,shape.image.content_type),"image"))
        return pages
    if suffix == ".pdf":
        pages=[]
        with fitz.open(path) as doc:
            if len(doc)>settings.max_pdf_pages:
                raise ValueError(f"This PDF has {len(doc)} pages; the configured limit is {settings.max_pdf_pages}.")
            for i,page in enumerate(doc,1):
                text=page.get_text()
                if text.strip():
                    pages.append((i,text,"text"))
                # Render pages with visual content (including scanned PDFs) for vision retrieval.
                if not settings.demo_mode and (page.get_images() or page.get_drawings() or not text.strip()):
                    pages.append((i,await describe_image(page.get_pixmap(matrix=fitz.Matrix(1,1)).tobytes("png"),"image/png"),"image"))
        return pages
    raise ValueError("Unsupported document format")

@traceable(name="document-ingestion")
async def ingest(document):
    with db() as c:
        row=c.get(DocumentRecord,document)
        if not row or row.status=="ready":
            return
        row.status='processing'
        row.error=None
    event(row.notebook,"ingestion",{"document":document,"status":"processing"})
    try:
        pages=await extract(Path(row.path))
        splitter=RecursiveCharacterTextSplitter(chunk_size=1200,chunk_overlap=200)
        chunks=[(page,text,modality) for page,body,modality in pages for text in splitter.split_text(body) if text.strip()]
        if not chunks:
            raise ValueError("No readable content. Scanned PDFs require live vision mode.")
        vectors=[]
        # Bound embedding request size for book-length documents.
        for start in range(0,len(chunks),64):
            vectors.extend(await embed([text for _,text,_ in chunks[start:start+64]]))
        with db() as c:
            if len(vectors)!=len(chunks):
                raise ValueError('Embedding endpoint returned an incomplete batch')
            c.execute(delete(ChunkRecord).where(ChunkRecord.document==document))
            c.add_all([ChunkRecord(id=uid(),document=document,page=p,text=t,modality=m,vector=v) for (p,t,m),v in zip(chunks,vectors)])
            c.get(DocumentRecord,document).status='ready'
        event(row.notebook,"ingestion",{"document":document,"status":"ready","chunks":len(chunks)})
    except Exception as exc:
        with db() as c:
            failed=c.get(DocumentRecord,document)
            failed.status='failed'
            failed.error=str(exc)[:500]
        event(row.notebook,"ingestion",{"document":document,"status":"failed"})
        raise
