import asyncio
import io
import json
import pytest
from sqlalchemy import select, func
from app.db_models import ChunkRecord, DocumentRecord, OutboxRecord
from fastapi.testclient import TestClient
from pptx import Presentation
from app.config import settings
from app.main import app
from app.auth import issue
from app.store import db
from app.ingestion import ingest
from app.rag import retrieve, rag_graph

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'data_dir',tmp_path)
    monkeypatch.setattr(settings,'demo_mode',True)
    with TestClient(app) as c:
        yield c

def headers(name='alice'):
    return {'Authorization':'Bearer '+issue(name)}

def notebook(client,owner='alice'):
    return client.post('/api/notebooks',headers=headers(owner),json={'title':'Distributed systems'}).json()['id']

def uploaded(client,nid,text='Kafka stores events durably. Consumers commit offsets after processing.'):
    response=client.post(f'/api/notebooks/{nid}/documents',headers=headers(),files={'file':('notes.md',text,'text/markdown')})
    assert response.status_code==202
    return response.json()['id']

async def test_upload_ingestion_chat_and_artifacts(client):
    nid=notebook(client)
    doc=uploaded(client,nid)
    with db() as c:
        assert c.scalar(select(OutboxRecord.sent))==0
        assert c.scalar(select(func.count()).select_from(ChunkRecord))==0
    await ingest(doc)
    await ingest(doc) # Redelivery is idempotent.
    with db() as c:
        assert c.scalar(select(func.count()).select_from(ChunkRecord))==1
    result=client.post(f'/api/notebooks/{nid}/chat',headers=headers(),json={'question':'How does Kafka work?'})
    frames=[json.loads(l[6:]) for l in result.text.splitlines() if l.startswith('data: ')]
    assert frames[-1]['type']=='done'
    assert any(e['type']=='token' and 'Kafka' in e['text'] for e in frames)
    assert any(e['type']=='sources' and e['sources'][0]['page']==1 for e in frames)
    quiz=client.post(f'/api/notebooks/{nid}/quiz',headers=headers(),json={'topic':'Kafka'})
    assert quiz.status_code==200
    assert quiz.json()['questions'][0]['citations']==[1]
    deck=client.post(f'/api/notebooks/{nid}/deck',headers=headers(),json={'topic':'Kafka'})
    assert deck.status_code==200
    artifact=deck.json()['id']
    file=client.get('/api/artifacts/'+artifact,headers=headers())
    prs=Presentation(io.BytesIO(file.content))
    assert len(prs.slides)==2
    assert 'Kafka' in prs.slides[0].shapes[1].text
    assert client.get('/api/artifacts/'+artifact,headers=headers('bob')).status_code==404

async def test_permission_isolation_on_every_surface(client):
    alice=notebook(client)
    bob=notebook(client,'bob')
    doc=uploaded(client,alice,'ALICE_PRIVATE_SECRET Kafka')
    await ingest(doc)
    assert await retrieve(bob,'bob','Kafka')==[]
    with pytest.raises(Exception) as exc:
        await retrieve(alice,'bob','Kafka')
    assert exc.value.status_code==404
    for path in ['documents','events']:
        assert client.get(f'/api/notebooks/{alice}/{path}',headers=headers('bob')).status_code==404
    for path,payload in [('chat',{'question':'secret'}),('quiz',{'topic':'secret'}),('deck',{'topic':'secret'})]:
        assert client.post(f'/api/notebooks/{alice}/{path}',headers=headers('bob'),json=payload).status_code==404
    assert client.post(f'/api/notebooks/{alice}/documents',headers=headers('bob'),files={'file':('a.md','bad')}).status_code==404
    assert len(client.get('/api/notebooks',headers=headers('bob')).json())==1

async def test_failed_ingestion_retry_and_validation(client):
    nid=notebook(client)
    response=client.post(f'/api/notebooks/{nid}/documents',headers=headers(),files={'file':('bad.exe',b'bad')})
    assert response.status_code==415
    assert client.post(f'/api/notebooks/{nid}/documents',headers=headers(),files={'file':('empty.md',b'')}).status_code==400
    doc=uploaded(client,nid,'\x00\xff')
    with db() as c:
        path=c.get(DocumentRecord,doc).path
    from pathlib import Path
    Path(path).write_bytes(b'\xff')
    with pytest.raises(UnicodeDecodeError):await ingest(doc)
    assert client.get(f'/api/notebooks/{nid}/documents',headers=headers()).json()[0]['status']=='failed'
    assert client.post(f'/api/notebooks/{nid}/documents/{doc}/retry',headers=headers()).status_code==202
    assert client.post(f'/api/notebooks/{nid}/documents/{doc}/retry',headers=headers()).status_code==409

async def test_bounded_adaptive_multihop(client,monkeypatch):
    from app import rag
    from langchain_core.messages import AIMessage
    nid=notebook(client);await ingest(uploaded(client,nid))
    monkeypatch.setattr(settings,'demo_mode',False)
    # Exercise actual graph branching with deterministic model and embedding substitutes.
    from app.models import demo_vector
    async def embeddings(texts):return [demo_vector(t) for t in texts]
    monkeypatch.setattr(rag,'embed',embeddings)
    class Model:
        async def ainvoke(self,messages):
            if 'Split the question' in messages[0][1]:return AIMessage(content='{"queries":["Kafka"]}')
            return AIMessage(content='{"sufficient":false,"queries":["offset commits"]}')
        async def astream(self,messages):
            yield AIMessage(content='Offsets are committed after processing [1].')
    monkeypatch.setattr(rag,'llm',lambda:Model())
    result=await rag_graph.ainvoke({'user':'alice','notebook':nid,'question':'How are offsets processed?','history':[]})
    assert result['hop']==3
    assert result['answer'].endswith('[1].')

def test_authentication(client):
    assert client.get('/api/notebooks').status_code in (401,403)
    assert client.get('/api/notebooks',headers={'Authorization':'Bearer invalid'}).status_code==401
    assert client.post('/api/login',json={'username':'alice','password':'bad'}).status_code==401
    assert client.post('/api/login',json={'username':'alice','password':settings.demo_password}).status_code==200

async def test_deep_agent_uses_scoped_search_tool(client,monkeypatch):
    from app import studio
    from langchain_core.language_models.fake_chat_models import FakeMessagesListChatModel
    from langchain_core.messages import AIMessage
    nid=notebook(client);await ingest(uploaded(client,nid))
    hits=await retrieve(nid,'alice','Kafka')
    monkeypatch.setattr(settings,'demo_mode',False)
    calls=[]
    async def scoped_search(notebook,user,query,limit=6):
        calls.append((notebook,user,query));return hits
    monkeypatch.setattr(studio,'retrieve',scoped_search)
    class ToolModel(FakeMessagesListChatModel):
        def bind_tools(self,tools,**kwargs):return self
    model=ToolModel(responses=[
        AIMessage(content='',tool_calls=[{'name':'search_sources','args':{'query':'Kafka offsets'},'id':'search1','type':'tool_call'}]),
        AIMessage(content=json.dumps({'questions':[{'question':'When are offsets committed?','answer':'After successful processing.','rubric':['Explains redelivery'],'citations':[1]}]}))])
    monkeypatch.setattr(studio,'llm',lambda:model)
    output,_=await studio.generate(nid,'alice','Kafka','quiz')
    assert output.questions[0].citations==[1]
    assert len(calls)==2
    assert all(n==nid and u=='alice' for n,u,q in calls)

async def test_image_description_retains_modality_and_page(client,monkeypatch):
    from app import ingestion
    from app.models import demo_vector
    from PIL import Image
    nid=notebook(client)
    buffer=io.BytesIO();Image.new('RGB',(20,20),'green').save(buffer,format='PNG')
    response=client.post(f'/api/notebooks/{nid}/documents',headers=headers(),files={'file':('diagram.png',buffer.getvalue(),'image/png')})
    doc=response.json()['id']
    with pytest.raises(ValueError,match='live mode'):await ingest(doc)
    monkeypatch.setattr(settings,'demo_mode',False)
    async def describe(raw,mime):return 'Diagram: Kafka producer sends an event to a broker.'
    async def embedding(texts):return [demo_vector(t) for t in texts]
    monkeypatch.setattr(ingestion,'describe_image',describe)
    monkeypatch.setattr(ingestion,'embed',embedding)
    await ingest(doc)
    with db() as c:
        row=c.scalar(select(ChunkRecord).where(ChunkRecord.document==doc))
        assert row.modality=='image' and row.page==1
        assert 'producer' in row.text

async def test_worker_failure_is_durable_before_commit(client,monkeypatch):
    from app import worker
    class Consumer:
        commits=0
        def __aiter__(self):return self
        async def __anext__(self):
            if self.commits:raise StopAsyncIteration
            from types import SimpleNamespace
            return SimpleNamespace(value=b'{"document":"bad"}')
        async def commit(self):self.commits+=1
    class Producer:
        messages=[]
        async def send_and_wait(self,topic,payload):self.messages.append((topic,json.loads(payload)))
    calls=[]
    async def fail(doc):calls.append(doc);raise ValueError('broken source')
    async def no_delay(seconds):pass
    monkeypatch.setattr(worker,'ingest',fail)
    monkeypatch.setattr(worker.asyncio,'sleep',no_delay)
    consumer,producer=Consumer(),Producer()
    await worker.consume(consumer,producer)
    assert len(calls)==3
    assert producer.messages==[('document.failed',{'document':'bad'})]
    assert consumer.commits==1

async def test_pdf_over_100_pages_and_configurable_limit(client,monkeypatch):
    import pymupdf
    from app.ingestion import extract
    path=settings.data_dir/'book.pdf'
    with pymupdf.open() as pdf:
        for i in range(101):
            page=pdf.new_page()
            page.insert_text((72,72),f'Learning material on page {i+1}')
        pdf.save(path)
    monkeypatch.setattr(settings,'max_pdf_pages',2000)
    pages=await extract(path)
    assert len(pages)==101 and pages[-1][0]==101
    assert 'page 101' in pages[-1][1]
    monkeypatch.setattr(settings,'max_pdf_pages',100)
    with pytest.raises(ValueError,match='101 pages; the configured limit is 100'):
        await extract(path)
